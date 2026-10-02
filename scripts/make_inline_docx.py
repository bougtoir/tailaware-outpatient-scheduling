"""Build manuscript_inline.docx: full text with figures embedded immediately
after the paragraph where they are first cited, and Tables 1-4 inserted at
their first-citation points (or appended if uncited)."""
import os, re, sys
import pandas as pd
from docx import Document
from docx.shared import Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAN = os.path.join(ROOT, "manuscript")
FIG = os.path.join(ROOT, "figures")
PROC = os.path.join(ROOT, "results", "processed")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx_math import (add_para, add_display_math, set_cell, fix_cell_text,
                       set_document_fonts, add_caption)

FIG_CAPTIONS = {
    "fig1_same_mean_tails": ("Figure 1.", "Same-mean ($E[S]=10$ min) service-time distributions with different upper tails (survival functions)."),
    "fig2_delay_propagation": ("Figure 2.", "Finite-session delay propagation under fixed mean-based slots ($x = E[S]$)."),
    "fig3_regret_surface": ("Figure 3.", "Fixed mean-based slot relative regret across families and CV, by session size $N$."),
    "fig4_misspec_matrix": ("Figure 4.", "True x assumed distribution matrix: excess cost of misspecification."),
    "fig5_pareto": ("Figure 5.", "Waiting vs idle+overtime trade-off under the uniform-interval sweep."),
    "fig6_cascade": ("Figure 6.", "Delay cascade from a single injected 45-min consultation."),
    "fig7_decision_map": ("Figure 7.", "Value-of-optimization map (bubble size = fixed-slot relative regret, $N=30$)."),
}

TABLES = [
    ("Table 1", "table1_design.csv", "Model and simulation design."),
    ("Table 2", "table2_policies.csv", "Scheduling policies and information requirements."),
    ("Table 3", "table3_performance.csv", "Representative policy performance and regret (focal specs)."),
    ("Table 4", "table4_complexity.csv", "Policy complexity vs benefit summary."),
]


def add_table(doc, csv_path, label, cap):
    df = pd.read_csv(csv_path)
    add_caption(doc, label + '.', cap)
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Light Grid Accent 1"
    for j, c in enumerate(df.columns):
        t.rows[0].cells[j].text = str(c)
    MAXR = 25
    for _, r in df.head(MAXR).iterrows():
        cells = t.add_row().cells
        for j, v in enumerate(r):
            s = f"{v:.4g}" if isinstance(v, float) else str(v)
            set_cell(cells[j], fix_cell_text(s))
    if len(df) > MAXR:
        doc.add_paragraph(f"(showing {MAXR} of {len(df)} rows; full data in results/processed/)")


def main():
    md = open(os.path.join(MAN, "manuscript.md")).read()
    doc = Document()
    set_document_fonts(doc)
    placed_figs = set()
    pending_tables = list(TABLES)
    para_lines = []

    def flush_para():
        if not para_lines:
            return
        b = " ".join(para_lines)
        para_lines.clear()
        add_para(doc, b)
        for tag, (label, cap) in FIG_CAPTIONS.items():
            mnum = re.search(r'S?(\d+)', tag[3:]).group(0)
            if tag in placed_figs:
                continue
            if re.search(rf"Fig(?:s)?\.\s*{mnum}\b", b):
                p = os.path.join(FIG, tag + ".png")
                if os.path.exists(p):
                    doc.add_picture(p, width=Inches(6))
                    add_caption(doc, label, cap)
                    placed_figs.add(tag)
        for t in list(pending_tables):
            if t[0] in b:
                add_table(doc, os.path.join(PROC, t[1]), t[0], t[2])
                pending_tables.remove(t)

    for block in md.split("\n"):
        b = block.rstrip()
        if not b:
            flush_para()
            continue
        if b.startswith("### "):
            flush_para()
            doc.add_heading(b[4:], level=2)
        elif b.startswith("## "):
            flush_para()
            doc.add_heading(b[3:], level=1)
        elif b.startswith("# "):
            flush_para()
            doc.add_heading(b[2:], level=0)
        elif b.startswith("$$") and b.endswith("$$"):
            flush_para()
            add_display_math(doc, b[2:-2])
        elif re.match(r"^\s*(?:[-*] |\d+\. )", b):
            flush_para()
            para_lines.append(b.strip())
        else:
            para_lines.append(b.strip())
    flush_para()
    for t in pending_tables:
        add_table(doc, os.path.join(PROC, t[1]), t[0], t[2])
    out = os.path.join(MAN, "manuscript_inline.docx")
    set_document_fonts(doc)
    doc.save(out)
    print("wrote", out, "figs placed:", sorted(placed_figs))


if __name__ == "__main__":
    main()
