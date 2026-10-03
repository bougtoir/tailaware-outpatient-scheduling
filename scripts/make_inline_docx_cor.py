#!/usr/bin/env python3
"""COR variant of make_inline_docx: reads manuscript/cor_manuscript.md,
places figures 1-8 (figures/) and 9-11 (figures_cor/) plus Tables 1-6
after first citation, writes cor_manuscript_inline.docx and
cor_manuscript.docx. Frozen Omega outputs untouched."""
import os, re, sys
import pandas as pd
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from make_inline_docx import (FIG_CAPTIONS, format_table, add_table, add_para,
                              add_display_math, set_document_fonts, add_caption)
from docx.shared import Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAN = os.path.join(ROOT, "manuscript")
FIG = os.path.join(ROOT, "figures")
FIG_COR = os.path.join(ROOT, "figures_cor")
PROC = os.path.join(ROOT, "results", "processed")
EXT = os.path.join(ROOT, "results", "ejor_extension")

FIG_CAPTIONS_COR = dict(FIG_CAPTIONS)
FIG_CAPTIONS_COR.update({
    "fig9_complexity_frontier": ("Figure 9.",
        "Complexity frontier: share of the uniform-to-full gap closed by $K$ piecewise-constant blocks ($N=30$, CV = 1.0, four families)."),
    "fig10_boundary_interior": ("Figure 10.",
        "(a) Fraction of the full nonuniform gap closed by six boundary (first/last three) vs six interior degrees of freedom ($N=30$). (b) Optimal interval minus uniform interval by session position."),
    "fig11_counterexamples_gap_vs_N": ("Figure 11.",
        "(a) Uniform-policy regret vs full design under nonstationary service, high overtime weight, and known heterogeneity. (b) Absolute gap and relative regret vs session size $N$."),
})

TABLES = [
    ("Table 1", "table1_design.csv", "Model and simulation design.", PROC),
    ("Table 2", "table2_policies.csv", "Scheduling policies and information requirements.", PROC),
    ("Table 3", "table3_performance.csv", "Representative policy performance and regret (focal specs).", PROC),
    ("Table 4", "table4_complexity.csv", "Policy complexity vs benefit summary.", PROC),
    ("Table 5", "table5_complexity_frontier.csv",
     "Share (%) of the uniform-to-full gap closed by $K$ block parameters ($N=30$).", EXT),
    ("Table 6", "table6_counterexamples.csv",
     "Counterexample regimes: expected cost of optimized uniform vs full nonuniform design.", EXT),
]

import make_inline_docx as mi
mi.TABLE_WIDTHS.update({
    "Table 5": [2.0, 0.57, 0.57, 0.57, 0.57, 0.57, 0.57, 0.57],
    "Table 6": [2.7, 1.2, 1.1, 1.2],
})


def fig_path(tag):
    for d in (FIG_COR, FIG):
        p = os.path.join(d, tag + ".png")
        if os.path.exists(p):
            return p
    return os.path.join(FIG, tag + ".png")


def main():
    md = open(os.path.join(MAN, "cor_manuscript.md")).read()
    doc = Document()
    set_document_fonts(doc)
    placed_figs = set()
    pending_tables = list(TABLES)
    para_lines = []

    def flush_para():
        if not para_lines:
            return
        b = " ".join(para_lines)
        para_lines.clear()
        add_para(doc, b)
        placements = []
        for tag, (label, cap) in FIG_CAPTIONS_COR.items():
            mnum = re.search(r'S?(\d+)', tag[3:]).group(0)
            if tag in placed_figs:
                continue
            match = re.search(rf"Fig(?:s)?\.\s*{mnum}\b", b)
            if match:
                placements.append((match.start(), "figure", (tag, label, cap)))
        for table in pending_tables:
            position = b.find(table[0])
            if position >= 0:
                placements.append((position, "table", table))
        for _, kind, item in sorted(placements):
            if kind == "figure":
                tag, label, cap = item
                doc.add_picture(fig_path(tag), width=Inches(6))
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
                add_caption(doc, label, cap)
                placed_figs.add(tag)
            else:
                label, csv_name, cap, d = item
                add_table(doc, os.path.join(d, csv_name), label, cap)
                pending_tables.remove(item)

    for block in md.split("\n"):
        b = block.rstrip()
        if not b:
            flush_para()
            continue
        if b.startswith("### "):
            flush_para()
            doc.add_heading(b[4:], level=2)
        elif b.startswith("## "):
            flush_para()
            doc.add_heading(b[3:], level=1)
        elif b.startswith("# "):
            flush_para()
            doc.add_heading(b[2:], level=0)
        elif b.startswith("$$") and b.endswith("$$"):
            flush_para()
            add_display_math(doc, b[2:-2])
        elif re.match(r"^\s*(?:[-*] |\d+\. )", b):
            flush_para()
            para_lines.append(b.strip())
        else:
            para_lines.append(b.strip())
    flush_para()
    assert not pending_tables, pending_tables
    assert placed_figs == set(FIG_CAPTIONS_COR), set(FIG_CAPTIONS_COR) - placed_figs
    set_document_fonts(doc)
    out = os.path.join(MAN, "cor_manuscript_inline.docx")
    doc.save(out)
    doc.save(os.path.join(MAN, "cor_manuscript.docx"))
    print("wrote", out, "figs placed:", sorted(placed_figs))


if __name__ == "__main__":
    main()
