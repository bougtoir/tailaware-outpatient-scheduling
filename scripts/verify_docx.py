"""Verify citation numbering, figure placement, and Word font properties."""
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from citations import cited_numbers

ROOT = Path(__file__).resolve().parents[1]
MAN = ROOT / "manuscript"


def verify_fonts(path):
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            root = etree.fromstring(archive.read(name))
            for fonts in root.iter(qn("w:rFonts")):
                assert not any("theme" in key.lower() for key in fonts.attrib), name
                for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                    assert fonts.get(qn(key)) == "Times New Roman", (name, key)
            for font in root.iter(qn("m:mathFont")):
                assert font.get(qn("m:val")) == "Times New Roman"
            for color in root.iter(qn("w:color")):
                assert dict(color.attrib) == {qn("w:val"): "000000"}, (name, color.attrib)
            for rpr in root.iter(qn("w:rPr")):
                color = rpr.find(qn("w:color"))
                assert color is not None, (name, "missing explicit text color")
            for fill in root.iter(qn("w14:textFill")):
                assert len(fill) == 1 and fill[0].tag == qn("a:srgbClr")
                assert fill[0].get("val") == "000000", name
            for scheme in root.iter(qn("a:fontScheme")):
                for font in scheme.iter():
                    if "typeface" in font.attrib:
                        assert font.get("typeface") == "Times New Roman"
    doc = Document(path)
    for style in doc.styles:
        size = style.element.find("./" + qn("w:rPr") + "/" + qn("w:sz"))
        expected = "36" if style.name in ("Title", "Title Char") else "24"
        assert size is not None and size.get(qn("w:val")) == expected, style.name
    for p in doc.element.body.iter(qn("w:p")):
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        assert not re.search(r"[$*]", text), text
        assert not re.search(r"[\x00-\x1f]", text), text
        for run in p.iter(qn("w:r")):
            size = run.find("./" + qn("w:rPr") + "/" + qn("w:sz"))
            assert size is not None
        if re.match(r"^(?:Figure|Table)\s+(?:S)?\d+\.", text):
            runs = p.findall(qn("w:r"))
            assert runs[0].find("./" + qn("w:rPr") + "/" + qn("w:b")) is not None
            for run in runs[1:]:
                assert run.find("./" + qn("w:rPr") + "/" + qn("w:b")) is None
            for run in list(p.iter(qn("w:r"))) + list(p.iter(qn("m:r"))):
                size = run.find("./" + qn("w:rPr") + "/" + qn("w:sz"))
                assert size.get(qn("w:val")) == "20", text
    return doc


def verify_references(doc):
    paragraphs = [p.text for p in doc.paragraphs]
    boundary = paragraphs.index("References")
    entries = paragraphs[boundary + 1:]
    labels = [int(re.match(r"^\[(\d+)\]", p).group(1)) for p in entries]
    cited = []
    for p in paragraphs[:boundary]:
        cited.extend(cited_numbers(p))
    assert labels == list(range(1, len(entries) + 1))
    assert list(dict.fromkeys(cited)) == labels
    assert set(cited) == set(labels)


def verify_figures(doc):
    nodes = list(doc.element.body)
    placed = []
    first = {}
    for i, node in enumerate(nodes):
        if node.tag != qn("w:p"):
            continue
        text = "".join(t.text or "" for t in node.iter(qn("w:t")))
        caption = re.match(r"^Figure\s+(\d+)\.", text)
        if caption:
            number = int(caption.group(1))
            placed.append(number)
            assert nodes[i - 1].find(".//" + qn("w:drawing")) is not None
            previous = i - 2
            while previous > first[number]:
                prior = nodes[previous]
                prior_text = "".join(t.text or "" for t in prior.iter(qn("w:t")))
                is_object = (
                    prior.tag == qn("w:tbl")
                    or prior.find(".//" + qn("w:drawing")) is not None
                    or re.match(r"^(?:Figure|Table)\s+", prior_text)
                )
                if not is_object:
                    break
                previous -= 1
            assert previous == first[number], (number, first[number], i)
            continue
        for match in re.finditer(r"Fig\.\s+(\d+)\b", text):
            first.setdefault(int(match.group(1)), i)
    assert placed == list(range(1, 8)), placed
    assert not any("Figure S" in p.text for p in doc.paragraphs)


def main():
    for name in ("manuscript.docx", "manuscript_inline.docx", "supplement.docx"):
        doc = verify_fonts(MAN / name)
        if name != "supplement.docx":
            verify_references(doc)
        if name == "manuscript_inline.docx":
            verify_figures(doc)
        print(name, "PASS")


if __name__ == "__main__":
    main()
