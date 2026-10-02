"""Shared helpers to add styled text (markdown bold/italic) and native Word
equations (OMML, converted from inline $...$ LaTeX) to python-docx paragraphs."""
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree
from latex2mathml.converter import convert as latex_to_mathml
from docx_equation import mathml_to_omml

TOKEN_RE = re.compile(r'(\*\*[^*]+\*\*|\$[^$\n]+?\$|\*[^*\n]+\*)')


def insert_omml(p, latex):
    mathml = latex_to_mathml(latex)
    p._element.append(mathml_to_omml(mathml))


def add_text(p, text, base_bold=False):
    """Append text to paragraph, honoring **bold**, *italic*, and $math$."""
    for tok in TOKEN_RE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = p.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith('$') and tok.endswith('$') and len(tok) > 2:
            insert_omml(p, tok[1:-1])
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = p.add_run(tok[1:-1])
            r.italic = True
            if base_bold:
                r.bold = True
        else:
            r = p.add_run(tok)
            if base_bold:
                r.bold = True


def add_para(doc, text, base_bold=False):
    p = doc.add_paragraph()
    add_text(p, text, base_bold=base_bold)
    return p


def add_display_math(doc, latex):
    """Standalone centered display equation."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    insert_omml(p, latex)
    return p


def set_cell(cell, text):
    """Set a table cell's text with inline-formatting/math support."""
    add_text(cell.paragraphs[0], text)


# presentation-only fixes applied to cell strings pulled from CSVs
CELL_FIXES = [
    ("D1=0; D_{i+1}=max(0, D_i + S_i - x_i)",
     "$D_1 = 0$; $D_{i+1} = \\max(0,\\, D_i + S_i - x_i)$"),
    ("Q85[S]", "$Q_{85}[S]$"),
    ("(c_w, c_i, c_o)", "$(c_w, c_i, c_o)$"),
    ("alpha=2.5,3.5", "$\\alpha = 2.5, 3.5$"),
    ("E[S]", "$E[S]$"),
]


def fix_cell_text(text):
    for old, new in CELL_FIXES:
        text = text.replace(old, new)
    return text


def set_document_fonts(doc, base_pt=12, title_pt=18):
    """Times New Roman everywhere; base_pt body text, larger title."""
    for style in doc.styles:
        size = title_pt if style.name in ("Title", "Title Char") else base_pt
        rpr = style.element.get_or_add_rPr()
        _set_run_properties(rpr, size)
        for nested in style.element.iter(qn("w:rPr")):
            _set_run_properties(nested, size)

    for part in doc.part.package.parts:
        if isinstance(part, XmlPart):
            for rf in part.element.iter(qn("w:rFonts")):
                _set_font_names(rf)
            for p in part.element.iter(qn("w:p")):
                style = p.find("./" + qn("w:pPr") + "/" + qn("w:pStyle"))
                size = title_pt if style is not None and style.get(qn("w:val")) == "Title" else base_pt
                for tag in ("w:r", "m:r"):
                    for run in p.iter(qn(tag)):
                        rpr = _run_properties(run)
                        existing = rpr.find(qn("w:sz"))
                        run_size = int(existing.get(qn("w:val"))) / 2 if existing is not None else size
                        _set_run_properties(rpr, run_size)
        elif part.content_type == "application/vnd.openxmlformats-officedocument.theme+xml":
            theme = parse_xml(part.blob)
            for tag in ("a:latin", "a:ea", "a:cs", "a:font"):
                for font in theme.iter(qn(tag)):
                    font.set("typeface", "Times New Roman")
            part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)
        elif str(part.partname).endswith(".xml"):
            root = parse_xml(part.blob)
            for rf in root.iter(qn("w:rFonts")):
                _set_font_names(rf)
            for style in root.iter(qn("w:style")):
                name = style.find(qn("w:name"))
                size = title_pt if name is not None and name.get(qn("w:val")) in ("Title", "Title Char") else base_pt
                for rpr in style.iter(qn("w:rPr")):
                    _set_run_properties(rpr, size)
            for defaults in root.iter(qn("w:rPrDefault")):
                for rpr in defaults.iter(qn("w:rPr")):
                    _set_run_properties(rpr, base_pt)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    defaults = doc.styles.element.find("./" + qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    if defaults is not None:
        _set_run_properties(defaults, base_pt)
    math_pr = doc.settings.element.find(qn("m:mathPr"))
    if math_pr is None:
        math_pr = OxmlElement("m:mathPr")
        doc.settings.element.append(math_pr)
    math_font = math_pr.find(qn("m:mathFont"))
    if math_font is None:
        math_font = OxmlElement("m:mathFont")
        math_pr.insert(0, math_font)
    math_font.set(qn("m:val"), "Times New Roman")


def _set_font_names(rf):
    for attr in list(rf.attrib):
        if "theme" in attr.lower():
            del rf.attrib[attr]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), "Times New Roman")


def _run_properties(run):
    rpr = run.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        math_pr = run.find(qn("m:rPr"))
        if math_pr is None:
            run.insert(0, rpr)
        else:
            math_pr.addnext(rpr)
    return rpr


def _set_run_properties(rpr, size):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    _set_font_names(rf)
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), str(int(size * 2)))


def add_caption(doc, label, text, size=10):
    """'Figure N.'/'Table N.' bold + caption normal, both at `size` pt."""
    p = doc.add_paragraph()
    r = p.add_run(label + ' ')
    r.bold = True
    r.font.size = Pt(size)
    # caption text may still contain $...$ math
    for tok in TOKEN_RE.split(text):
        if not tok:
            continue
        if tok.startswith('$') and tok.endswith('$') and len(tok) > 2:
            insert_omml(p, tok[1:-1])
            _size_omml(p, size)
        else:
            rr = p.add_run(tok)
            rr.font.size = Pt(size)
    return p


def _size_omml(p, pt):
    for mr in p._element.findall('.//' + qn('m:r')):
        _set_run_properties(_run_properties(mr), pt)
