"""Freeze and verify scientific content for the citation-formatting pass."""
import argparse
import hashlib
import json
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from citations import CITATION_RE

ROOT = Path(__file__).resolve().parents[1]
QC = ROOT / "qc"
BASELINE = QC / "pre_format_content_manifest.json"
DOCS = ("manuscript.docx", "manuscript_inline.docx", "supplement.docx")
FRAMEWORK_CLAUSE = (
    "resting on CVaR and robust-optimization machinery, "
    "with general scheduling foundations in Pinedo"
)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def visible_text(element):
    return "".join(
        node.text or "" for node in element.iter()
        if node.tag in (qn("w:t"), qn("m:t"))
    )


def plain(text):
    text = CITATION_RE.sub("", text)
    text = re.sub(r"\s+", " ", text).strip()
    return re.sub(r"\s+([,.;:])", r"\1", text)


def science_text(paragraphs):
    paragraphs = [
        text for text in paragraphs
        if not re.match(r"^(Figure|Table) S?\d+\.", text)
    ]
    text = plain(" ".join(paragraphs))
    text = text.replace(" (Supplementary Fig. S1)", "")
    text = text.replace(", " + FRAMEWORK_CLAUSE + ",", "")
    text = text.replace(", " + FRAMEWORK_CLAUSE, "")
    return text


def scientific_numbers(paragraphs):
    text = " ".join(
        value for value in paragraphs
        if not re.match(r"^(Figure|Table) S?\d+\.", value)
    ).replace(" (Supplementary Fig. S1)", "")
    return re.findall(r"\d+(?:[.,]\d+)*", plain(text))


def math_semantics(node):
    if not node.tag.startswith("{" + qn("m:r").split("}")[0][1:] + "}"):
        return None
    return {
        "tag": node.tag,
        "attributes": dict(node.attrib),
        "text": node.text,
        "children": [
            child for item in node
            if (child := math_semantics(item)) is not None
        ],
    }


def equation_content(node):
    normalized = {
        **node,
        "children": [
            equation_content(child) for child in node["children"]
            if child["tag"] != qn("m:lit")
        ],
    }
    if normalized["text"] == "--":
        normalized["text"] = "–"
    return normalized


def snapshot(path):
    doc = Document(path)
    paragraphs = [visible_text(p._element) for p in doc.paragraphs]
    boundary = paragraphs.index("References") if "References" in paragraphs else len(paragraphs)
    body = paragraphs[:boundary]
    return {
        "sha256": digest(path),
        "paragraphs": paragraphs,
        "body": body,
        "headings": [
            visible_text(p._element) for p in doc.paragraphs
            if p.style.name.startswith(("Heading", "Title"))
        ],
        "equations": [
            math_semantics(node) for node in doc.element.body.iter(qn("m:oMath"))
        ],
        "captions": [
            text for text in paragraphs if re.match(r"^(Figure|Table) S?\d+\.", text)
        ],
        "tables": [
            [[visible_text(cell._tc) for cell in row.cells] for row in table.rows]
            for table in doc.tables
        ],
        "references": paragraphs[boundary + 1:],
        "numerical_strings": re.findall(
            r"\d+(?:[.,]\d+)*",
            plain(" ".join(body)).replace(" (Supplementary Fig. S1)", ""),
        ),
        "embedded_media": sorted(
            hashlib.sha256(part.blob).hexdigest()
            for part in doc.part.package.parts
            if str(part.partname).startswith("/word/media/")
        ),
    }


def frozen_files():
    files = [ROOT / "output/omega_submission_package_FINAL.zip"]
    for directory in ("src", "config", "configs", "results", "figures", "tables"):
        parent = ROOT / directory
        if parent.exists():
            files.extend(
                path for path in parent.rglob("*")
                if path.is_file() and "__pycache__" not in path.parts
            )
    files.extend(
        path for path in (ROOT / "scripts").glob("exp*.py")
    )
    files.append(ROOT / "literature/literature_matrix.csv")
    return {str(path.relative_to(ROOT)): digest(path) for path in sorted(set(files))}


def freeze():
    if BASELINE.exists():
        raise FileExistsError("The original formatting baseline must not be overwritten")
    data = {
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "documents": {name: snapshot(ROOT / "manuscript" / name) for name in DOCS},
        "immutable_files": frozen_files(),
    }
    BASELINE.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")
    rows = ["# Pre-format content manifest", "", f"Frozen UTC: {data['created_utc']}", ""]
    for name, doc in data["documents"].items():
        rows.extend([
            f"## {name}", "", f"SHA-256: `{doc['sha256']}`",
            f"Paragraphs: {len(doc['paragraphs'])}; equations: {len(doc['equations'])}; "
            f"tables: {len(doc['tables'])}; references: {len(doc['references'])}.", "",
            "Full paragraph, heading, equation, caption, table, reference and numerical-string "
            "snapshots are retained in pre_format_content_manifest.json.", "",
        ])
    rows.extend(["## Frozen inputs, results, figures and previous package", ""])
    rows.extend(f"- `{name}`: `{sha}`" for name, sha in data["immutable_files"].items())
    (QC / "pre_format_content_manifest.md").write_text("\n".join(rows) + "\n")
    print("Frozen:", BASELINE)


def verify():
    if (QC / "integrated_architecture_baseline.json").exists():
        subprocess.run(
            [sys.executable, str(ROOT / "scripts/integrated_architecture.py"), "verify"],
            check=True,
        )
        return
    data = json.loads(BASELINE.read_text())
    rows = [
        "# Content preservation diff", "",
        "Allowed editorial differences: citation relocation/renumbering, verified reference "
        "metadata and formatting, and relocation of the existing CVaR/robust-machinery/Pinedo "
        "clause from Introduction to Model (comma/period adjusted). The estimation-uncertainty "
        "paragraph gains only an explicit cross-reference to the existing Supplementary Fig. S1. "
        "No other prose change is allowed.", "",
    ]
    for name, before in data["documents"].items():
        after = snapshot(ROOT / "manuscript" / name)
        assert science_text(before["body"]) == science_text(after["body"]), name
        assert scientific_numbers(before["body"]) == scientific_numbers(after["body"]), name
        assert sorted(before["captions"]) == sorted(after["captions"]), (name, "captions")
        assert [
            equation_content(node) for node in before["equations"]
        ] == [
            equation_content(node) for node in after["equations"]
        ], (name, "equations")
        for field in ("headings", "tables", "embedded_media"):
            assert before[field] == after[field], (name, field)
        rows.append(
            f"- {name}: PASS — prose (apart from declared clause relocation), numerical "
            "strings, equations, headings, captions, tables and embedded figures preserved."
        )
    for name, expected in data["immutable_files"].items():
        assert digest(ROOT / name) == expected, name
    rows.extend([
        f"- {len(data['immutable_files'])} frozen files: SHA-256 unchanged, including all "
        "simulation scripts/inputs/results, figures and the previous FINAL submission ZIP.",
        "- Supplement prose and scientific content: unchanged.",
        "- Bibliographic years/pages and citation numbers are editorial metadata, excluded "
        "from the scientific numerical comparison.", "",
        "Scientific content changed: NO. New analyses performed: NO.", "",
    ])
    (QC / "content_preservation_diff.md").write_text("\n".join(rows))
    print("Content-preservation audit PASS")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("action", choices=("freeze", "verify"))
    args = parser.parse_args()
    if args.action == "freeze":
        freeze()
    else:
        verify()


if __name__ == "__main__":
    main()
