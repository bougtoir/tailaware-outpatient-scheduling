#!/usr/bin/env python3
"""Build the SEPS (Socio-Economic Planning Sciences) manuscript variant.

Wraps the COR manuscript template (which already injects the frozen
extension numbers) and applies SEPS-specific framing edits only:
service-sector planning literature anchors, planning-decision framing in
Abstract/Discussion, and the SEPS reference set. Frozen scientific content,
numbers, figures and tables are unchanged. Writes
manuscript/seps_manuscript.{md,docx}; COR and Omega artifacts untouched.
"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_manuscript as bm
import build_manuscript_cor as bmc

ROOT = bm.ROOT
MAN = os.path.join(ROOT, "manuscript")


def number_references_seps(text):
    path = os.path.join(ROOT, "literature", "normalized_references_seps.json")
    with open(path) as stream:
        records = json.load(stream)
    entries = {}
    dois = set()
    for record in records:
        key = record["citation_key"]
        if key in entries:
            raise ValueError(f"Duplicate reference key: {key}")
        if record["doi"] in dois:
            raise ValueError(f"Duplicate DOI: {record['doi']}")
        dois.add(record["doi"])
        entries[key] = record["formatted_entry"]
    order = {}

    def replace(match):
        labels = []
        for key in match.group(1).split(";"):
            if key not in entries:
                raise ValueError(f"Unknown citation key: {key}")
            if key not in order:
                order[key] = len(order) + 1
            labels.append(order[key])
        return bm.format_citation(labels)

    text = re.sub(r"\[@([a-z0-9;]+)\]", replace, text)
    uncited = entries.keys() - order.keys()
    if uncited:
        raise ValueError(f"Uncited references: {sorted(uncited)}")
    return text + "\n\n".join(f"[{n}] {entries[key]}" for key, n in order.items()) + "\n"


def build_template_seps():
    t = bmc.build_template()

    # Abstract: rewrite for the SEPS planning audience and the ~250-word
    # unstructured-abstract norm (previous journal version was ~315 words).
    old_abs_start = t.find("Outpatient appointment systems are typically designed")
    old_abs_end = t.find("under which added complexity pays.")
    assert old_abs_start > 0 and old_abs_end > old_abs_start
    new_abs = (
        "Outpatient appointment systems commit session capacity before "
        "demand is realized, typically dividing session length by mean "
        "consultation time — yet consultation durations are right-skewed and "
        "their upper tails vary across clinics. We ask when fixed, "
        "mean-based slots fail under service-time distributional uncertainty, "
        "and how much scheduling sophistication is worth buying. Using a "
        "finite-horizon Lindley-type delay recursion evaluated by large-scale "
        "Monte Carlo simulation, we hold the mean fixed and vary only family "
        "and tail shape across gamma, Weibull, lognormal, mixture, and "
        "Pareto-type designs. Fixed slots incur "
        "{fm_regret_min}–{fm_regret_max} relative regret "
        "versus an optimized benchmark at session size $N = 30$ (scenario "
        "mean {fm_rel_N0} at $N = 20$, rising to {fm_rel_N2} at $N = 50$). "
        "A single "
        "SAA-optimized uniform interval recovers nearly all attainable "
        "benefit (maximum regret {uni_regret_max} versus oracle): ten "
        "contiguous blocks over the twenty-nine interval positions capture "
        "{cx_k10_min}–{cx_k10_max}% of the "
        "residual gap, concentrated "
        "in boundary positions. Misspecifying the family can erase these "
        "gains — under-designing costs about three times over-designing — "
        "and under a genuinely heavy-tailed truth, moment-fitted designs "
        "remain unreliable even at 1,000 historical observations. "
        "Constructed counterexamples mark where complexity pays: "
        "nonstationarity, structured heterogeneity, dominant overtime cost. "
        "For service planners the output is a decision rule — how much "
        "scheduling complexity to buy, and which distributional measurements "
        "justify it — rather than a new algorithm.")
    t = t[:old_abs_start] + new_abs + t[old_abs_end + len("under which added complexity pays."):]

    # Introduction: SEPS scholarly-conversation anchors (replacing the
    # COR-journal anchors).
    old_lit = ("Queueing and simulation models\n"
               "of appointment-driven systems are established in this journal\n"
               "[@creemers2008;dogru2023], along with quantile-objective\n"
               "designs [@begencao2021] and multi-stage scheduling under\n"
               "limited distributional information [@zhouyue2021].")
    assert old_lit in t
    t = t.replace(old_lit,
                  "Queueing and simulation models\n"
                  "of appointment-driven systems are established tools in this\n"
                  "journal's service-planning literature\n"
                  "[@shoaib2023;alaouchiche2024], and slot-structure design for\n"
                  "clinical scheduling under patient variability has itself\n"
                  "been studied here [@chakraborty2013], alongside stochastic\n"
                  "planning of elective-care capacity [@acuna2024].")

    # Keywords: add service-planning descriptor.
    old_kw = "decision rules"
    assert old_kw in t
    t = t.replace(old_kw, "service planning; decision rules", 1)

    # Discussion: connect to the SEPS planning conversation after the
    # interval-optimization comparison paragraph.
    old_disc = ("the boundary positions where that value\n"
                "concentrates.\n")
    assert old_disc in t
    t = t.replace(old_disc,
                  "the boundary positions where that value\n"
                  "concentrates. Within this journal, slot-structure design\n"
                  "for clinical scheduling [@chakraborty2013],\n"
                  "simulation-based healthcare planning\n"
                  "[@shoaib2023;alaouchiche2024;acuna2024], and patient-flow\n"
                  "studies linking accumulated delay to provider burden\n"
                  "[@cildoz2020] form the service-planning conversation that\n"
                  "our complexity analysis extends: the planning question is\n"
                  "not only which interval to use but how much scheduling\n"
                  "complexity is worth buying.\n")

    # Discussion: drop the COR-specific DRO anchor (zhouyue2021 was a COR
    # journal addition).
    old_dro = "model risk [@kong2013;mak2015;rahimian2022;zhouyue2021]."
    assert old_dro in t
    t = t.replace(old_dro, "model risk [@kong2013;mak2015;rahimian2022].")

    # Limitations: replace the COR interruption reference with the
    # already-cited interruption/lateness work.
    old_int = "Interruption-aware designs [@dogru2023] address a complementary realism axis."
    assert old_int in t
    t = t.replace(old_int,
                  "Interruption- and lateness-aware designs [@klassen2013] "
                  "address a complementary realism axis.")
    # Precision: K indexes contiguous blocks, not parameters.
    old_k = ("$K = 10$\nof the 29 adjustable parameters captures")
    assert old_k in t
    t = t.replace(old_k, "$K = 10$ blocks\nover the 29 adjustable intervals captures")

    # Hostile-review repair R1: Zhou et al. (2021) is finite-horizon bounds
    # plus numerical near-optimality of the constant policy with a
    # "dome-shaped" optimal vector, not merely asymptotic theory. Credit it
    # explicitly and reposition the increment.
    old_zhou = ("and asymptotic theory shows constant or critical-load\n"
                "schedules to be optimal in fluid and diffusion limits\n"
                "[@armony2019;zhou2021]. What finite sessions gain beyond a\n"
                "single optimized scalar interval is the empirical question\n"
                "we quantify.")
    assert old_zhou in t
    t = t.replace(old_zhou,
                  "and the constant (uniform) job-allowance policy itself is\n"
                  "known to be near-optimal — asymptotically in fluid and\n"
                  "diffusion limits [@armony2019] and, with explicit\n"
                  "finite-horizon performance bounds and numerical evidence,\n"
                  "even for small sessions [@zhou2021]; the optimal interval\n"
                  "vector in that literature varies mainly at the beginning\n"
                  "and end of the schedule. What those results do not provide\n"
                  "is a quantified map from distributional family — and from\n"
                  "misspecification of that family — to the finite-session\n"
                  "value of added scheduling complexity under an\n"
                  "overtime-inclusive cost. That is the empirical question we\n"
                  "quantify.")

    # R1 (cont.): Contribution 3 frames the simplicity result as
    # corroboration + quantification, not discovery.
    old_c3 = ("3. **A simplicity result.** A single scalar, SAA-optimized uniform interval\n"
              "   captures nearly all attainable benefit; nonuniform, CVaR-aware and\n"
              "   distributionally robust policies add little incremental value.")
    assert old_c3 in t
    t = t.replace(old_c3,
                  "3. **A simplicity result, quantified.** A single scalar,\n"
                  "   SAA-optimized uniform interval captures nearly all\n"
                  "   attainable benefit — a finite-session, overtime-inclusive\n"
                  "   corroboration and quantification of constant-policy\n"
                  "   near-optimality theory; nonuniform, CVaR-aware and\n"
                  "   distributionally robust policies add little incremental\n"
                  "   value.")

    # R1 (cont.): Contribution 4 acknowledges the known boundary structure.
    old_c4 = ("4. **Value of schedule complexity.** A complexity frontier over\n"
              "   block-structured interval families shows the nonuniform\n"
              "   benefit concentrates in a few boundary degrees of freedom;\n"
              "   interior slot-by-slot freedom adds almost nothing.")
    assert old_c4 in t
    t = t.replace(old_c4,
                  "4. **Value of schedule complexity.** A complexity frontier\n"
                  "   over block-structured interval families quantifies the\n"
                  "   residual nonuniform benefit and locates it in a few\n"
                  "   boundary degrees of freedom — consistent with the\n"
                  "   boundary-varied optimal vectors known in that\n"
                  "   literature; interior slot-by-slot freedom adds almost\n"
                  "   nothing.")

    # R1 (cont.): Discussion comparison sentence credits finite-horizon
    # bounds and the dome-shaped precedent explicitly.
    old_d = ("and the asymptotic optimality of constant policies\n"
             "[@armony2019;zhou2021] by quantifying finite-session regret —\n"
             "with overtime included — the marginal value of added schedule\n"
             "complexity, and the boundary positions where that value\n"
             "concentrates.")
    assert old_d in t
    t = t.replace(old_d,
                  "and the finite-horizon near-optimality of constant\n"
                  "job-allowance policies [@armony2019;zhou2021] by\n"
                  "quantifying finite-session regret — with overtime\n"
                  "included — the marginal value of added schedule\n"
                  "complexity, and the boundary positions where that value\n"
                  "concentrates — the positions whose importance the\n"
                  "dome-shaped optimal schedules of that literature\n"
                  "anticipate.")

    # R3: symmetric statistic labels in the estimation sentence.
    old_est = ("median estimation regret is already small\n"
               "at $n = 250$ ({est250:.2f} cost units; Q90 {est250q:.2f}) versus $n = 50$\n"
               "({est50:.2f})")
    assert old_est in t
    t = t.replace(old_est,
                  "median estimation regret is already small\n"
                  "at $n = 250$ (median {est250:.2f} cost units; median Q90\n"
                  "{est250q:.2f}) versus $n = 50$ (median Q90 {est50:.2f})")

    # R4: bound is 2.51% rendered 2.5 — qualify with "about".
    old_ws = "maximum regret stays at or below {ws_uni_max}\nacross the weight grid."
    assert old_ws in t
    t = t.replace(old_ws,
                  "maximum regret stays at or below about {ws_uni_max}\n"
                  "across the weight grid.")

    # Final-submission F3: the "oracle" is a K-block SAA benchmark, not a
    # certified global optimum — qualify it at first use and label it SAA
    # where it appears unqualified in abstract, highlights, contribution 2.
    old_or = ("The\noracle $\\pi^*$ optimizes a piecewise-constant interval vector under the\n"
              "true distribution by sample average approximation (SAA); regret is")
    assert old_or in t
    t = t.replace(old_or,
                  "The\noracle $\\pi^*$ optimizes a piecewise-constant interval vector under the\n"
                  "true distribution by sample average approximation (SAA) —\n"
                  "a computational best-known schedule within that block\n"
                  "family, not a certified global optimum; regret is")
    old_ab_or = "benefit (maximum regret {uni_regret_max} versus oracle): ten "
    assert old_ab_or in t
    t = t.replace(old_ab_or, "benefit (maximum regret {uni_regret_max} versus the SAA oracle): ten ")
    old_hl_or = "caps regret at {uni_regret_max} vs oracle"
    assert old_hl_or in t
    t = t.replace(old_hl_or, "caps regret at {uni_regret_max} vs the SAA oracle")
    old_c2_or = "Against a common-random-number oracle we show"
    assert old_c2_or in t
    t = t.replace(old_c2_or, "Against a common-random-number SAA oracle we show")
    return t


def main():
    n = bm.numbers()
    pct = bm.pct
    kw = dict(
        title="When Does Appointment-Scheduling Complexity Pay? "
              "Service-Time Uncertainty and the Value of Simple Planning Rules",
        fm_regret_min=pct(n["fm_regret_rel_min"]),
        fm_regret_max=pct(n["fm_regret_rel_max"]),
        uni_regret_max=pct(n["uni_regret_rel_max"]),
        fm_wait_ln15=n["fm_wait_ln15"],
        fm_wait30_ln15=pct(n["fm_wait30_ln15"]),
        pareto_fm_rel=pct(n["pareto_a25_fm_rel"]),
        mis_max=n["mis_max_excess"],
        est250=n["est_n250_regret_q50_med"],
        fm_rel_N0=pct(n["fm_rel_N20"]), fm_rel_N1=pct(n["fm_rel_N30"]),
        fm_rel_N2=pct(n["fm_rel_N50"]),
        fm_abs_cv0=n["fm_abs_cv025"], fm_abs_cv1=n["fm_abs_cv200"],
        pareto_est_q90=n["pareto_est_q90_n1000"],
        mis_under=n["mis_under"], mis_over=n["mis_over"],
        uni_mean=pct(n["uni_mean_regret"]),
        nonunif_inc_min=n["nonunif_inc_min"], nonunif_inc_max=n["nonunif_inc_max"],
        ws_uni_max=pct(n["ws_uni_max"]),
        rob_frac_ns=pct(n["rob_tail_frac_ns"]),
        est250q=n["est_n250_regret_q90_med"],
        est50=n["est_n50_regret_q90_med"],
        cascade_len=n["cascade_len_mid"],
        cascade_cum=n["cascade_cum_mid"],
        rob_frac=pct(n["rob_tail_better_frac"]),
    )
    kw.update(bmc.extension_numbers())
    text = build_template_seps().format(**kw)
    text = bm.unwrap_paragraphs(number_references_seps(text))
    out_md = os.path.join(MAN, "seps_manuscript.md")
    with open(out_md, "w") as f:
        f.write(text)
    print("wrote", out_md)

    from docx import Document
    from docx_math import add_para, add_display_math, set_document_fonts
    doc = Document()
    set_document_fonts(doc)
    para_lines = []

    def flush_para():
        if para_lines:
            add_para(doc, " ".join(para_lines))
            para_lines.clear()

    for block in text.split("\n"):
        b = block.rstrip()
        if not b:
            flush_para()
            continue
        if b.startswith("## "):
            flush_para()
            doc.add_heading(b[3:], level=1)
        elif b.startswith("### "):
            flush_para()
            doc.add_heading(b[4:], level=2)
        elif b.startswith("# "):
            flush_para()
            doc.add_heading(b[2:], level=0)
        elif b.startswith("$$") and b.endswith("$$"):
            flush_para()
            add_display_math(doc, b[2:-2])
        elif re.match(r"^\s*(?:[-*] |\d+\. )", b):
            flush_para()
            para_lines.append(b.strip())
        else:
            para_lines.append(b.strip())
    flush_para()
    set_document_fonts(doc)
    out = os.path.join(MAN, "seps_manuscript.docx")
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
